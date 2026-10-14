import io
import csv
import pytest
import openpyxl
from httpx import AsyncClient, ASGITransport
from app.main import app
from app.core.security import create_access_token, TokenPayload
from app.core.database import connect_to_mongo, get_tenant_db

TENANT_ID = "65f2a1b9a000000000000001"

@pytest.fixture(autouse=True)
async def init_db():
    try:
        await connect_to_mongo()
    except Exception:
        pass
    yield

@pytest.fixture
def auth_token():
    payload = TokenPayload(
        sub="usr_test_admin",
        email="admin@quickbill.local",
        roles=["tenant_admin"],
        permissions=["all"],
        default_business_id=TENANT_ID,
        authorized_business_ids=[TENANT_ID]
    )
    return create_access_token(payload)

@pytest.mark.asyncio
async def test_download_templates(auth_token):
    headers = {"Authorization": f"Bearer {auth_token}"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Download CSV template
        res_csv = await client.get("/api/v1/items/template?format=csv", headers=headers)
        assert res_csv.status_code == 200
        assert "text/csv" in res_csv.headers.get("content-type", "")
        csv_text = res_csv.text
        assert "Product Name" in csv_text
        assert "Sale Price" in csv_text
        assert "Category" in csv_text

        # 2. Download Excel template
        res_xlsx = await client.get("/api/v1/items/template?format=xlsx", headers=headers)
        assert res_xlsx.status_code == 200
        assert "spreadsheetml" in res_xlsx.headers.get("content-type", "")
        
        # Verify valid openpyxl workbook
        wb = openpyxl.load_workbook(io.BytesIO(res_xlsx.content))
        assert "Products" in wb.sheetnames
        assert "Instructions & Branches" in wb.sheetnames
        ws = wb["Products"]
        headers_row = [cell.value for cell in ws[1]]
        assert "Product Name" in headers_row
        assert "Sale Price" in headers_row


@pytest.mark.asyncio
async def test_import_csv_new_products_and_auto_categories(auth_token):
    headers = {"Authorization": f"Bearer {auth_token}"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Prepare CSV data with new category "Gourmet Bakery 2026"
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow(["Product Name", "Category", "Barcode", "SKU", "Sale Price", "MRP", "Purchase Price", "Opening Stock", "Unit"])
        writer.writerow(["Artisan Sourdough Loaf", "Gourmet Bakery 2026", "890998877001", "SKU-BAK-001", "120.00", "140.00", "70.00", "15.0", "loaf"])
        writer.writerow(["Almond Croissant", "Gourmet Bakery 2026", "890998877002", "SKU-BAK-002", "85.00", "95.00", "45.00", "20.0", "pcs"])
        csv_bytes = csv_buf.getvalue().encode("utf-8")

        files = {"file": ("products.csv", csv_bytes, "text/csv")}
        data = {
            "update_existing": "true",
            "create_categories": "true"
        }

        res = await client.post("/api/v1/items/import", headers=headers, files=files, data=data)
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["success"] is True
        assert body["createdCount"] == 2
        assert body["errorsCount"] == 0
        assert "Gourmet Bakery 2026" in body["categoriesCreated"]

        # Verify items exist in catalog
        list_res = await client.get("/api/v1/items?search=Sourdough", headers=headers)
        assert list_res.status_code == 200
        items = list_res.json()["data"]
        assert len(items) >= 1
        item = items[0]
        assert item["name"] == "Artisan Sourdough Loaf"
        assert item["category"] == "Gourmet Bakery 2026"
        assert float(item["currentStock"]) == 15.0


@pytest.mark.asyncio
async def test_import_update_existing_by_barcode_and_sku(auth_token):
    headers = {"Authorization": f"Bearer {auth_token}"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. First import
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow(["Product Name", "Category", "Barcode", "SKU", "Sale Price", "MRP", "Purchase Price", "Opening Stock"])
        writer.writerow(["Premium Olive Oil 500ml", "Edible Oils", "890776655001", "SKU-OIL-999", "350.00", "390.00", "280.00", "10.0"])
        
        files1 = {"file": ("oil.csv", csv_buf.getvalue().encode("utf-8"), "text/csv")}
        res1 = await client.post("/api/v1/items/import", headers=headers, files=files1, data={"update_existing": "true", "create_categories": "true"})
        assert res1.status_code == 200
        assert res1.json()["createdCount"] == 1

        # 2. Second import updating price and adding stock for the same barcode
        csv_buf2 = io.StringIO()
        writer2 = csv.writer(csv_buf2)
        writer2.writerow(["Product Name", "Category", "Barcode", "SKU", "Sale Price", "MRP", "Purchase Price", "Opening Stock"])
        writer2.writerow(["Premium Extra Virgin Olive Oil 500ml", "Edible Oils", "890776655001", "SKU-OIL-999", "380.00", "420.00", "300.00", "25.0"])

        files2 = {"file": ("oil_update.csv", csv_buf2.getvalue().encode("utf-8"), "text/csv")}
        res2 = await client.post("/api/v1/items/import", headers=headers, files=files2, data={"update_existing": "true", "create_categories": "true"})
        assert res2.status_code == 200
        body2 = res2.json()
        assert body2["updatedCount"] == 1
        assert body2["createdCount"] == 0

        # Verify updated details in items API
        list_res = await client.get("/api/v1/items?search=890776655001", headers=headers)
        assert list_res.status_code == 200
        item = list_res.json()["data"][0]
        assert item["name"] == "Premium Extra Virgin Olive Oil 500ml"
        assert float(item["salePrice"]) == 380.00
        assert float(item["currentStock"]) == 25.0


@pytest.mark.asyncio
async def test_import_skip_existing_when_disabled(auth_token):
    headers = {"Authorization": f"Bearer {auth_token}"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Create initial item
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow(["Product Name", "Barcode", "SKU", "Sale Price"])
        writer.writerow(["Test Skippable Item", "890112233445", "SKU-SKIP-01", "100.00"])
        
        files1 = {"file": ("skip.csv", csv_buf.getvalue().encode("utf-8"), "text/csv")}
        await client.post("/api/v1/items/import", headers=headers, files=files1, data={"update_existing": "true", "create_categories": "true"})

        # Upload again with update_existing = false
        files2 = {"file": ("skip.csv", csv_buf.getvalue().encode("utf-8"), "text/csv")}
        res2 = await client.post("/api/v1/items/import", headers=headers, files=files2, data={"update_existing": "false", "create_categories": "true"})
        assert res2.status_code == 200
        body2 = res2.json()
        assert body2["skippedCount"] == 1
        assert body2["createdCount"] == 0
        assert body2["updatedCount"] == 0


@pytest.mark.asyncio
async def test_import_excel_xlsx_file(auth_token):
    headers = {"Authorization": f"Bearer {auth_token}"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Generate an in-memory xlsx workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Products"
        ws.append(["Product Name", "Category", "Barcode", "SKU", "Sale Price", "MRP", "Purchase Price", "Opening Stock", "Unit"])
        ws.append(["Dark Chocolate Bar 85%", "Sweets & Confectionery", "890554433221", "SKU-CHOC-85", 150.00, 160.00, 95.00, 30.0, "bar"])
        
        excel_buf = io.BytesIO()
        wb.save(excel_buf)
        excel_bytes = excel_buf.getvalue()

        files = {"file": ("chocolates.xlsx", excel_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        res = await client.post("/api/v1/items/import", headers=headers, files=files, data={"update_existing": "true", "create_categories": "true"})
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["createdCount"] == 1
        assert "Sweets & Confectionery" in body["categoriesCreated"]


@pytest.mark.asyncio
async def test_import_row_error_reporting(auth_token):
    headers = {"Authorization": f"Bearer {auth_token}"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow(["Product Name", "Sale Price"])
        writer.writerow(["", "50.00"]) # Row 2: Missing name
        writer.writerow(["Valid Item X", "INVALID_PRICE"]) # Row 3: Invalid price
        writer.writerow(["Valid Item Y", "75.00"]) # Row 4: Valid

        files = {"file": ("errors.csv", csv_buf.getvalue().encode("utf-8"), "text/csv")}
        res = await client.post("/api/v1/items/import", headers=headers, files=files, data={"update_existing": "true", "create_categories": "true"})
        assert res.status_code == 200
        body = res.json()
        assert body["createdCount"] == 1
        assert body["errorsCount"] == 2
        err_rows = [e["row"] for e in body["errors"]]
        assert 2 in err_rows
        assert 3 in err_rows


@pytest.mark.asyncio
async def test_import_products_with_location_reflected_in_catalog(auth_token):
    headers = {"Authorization": f"Bearer {auth_token}"}
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Fetch tenant locations
        loc_res = await client.get("/api/v1/locations", headers=headers)
        assert loc_res.status_code == 200
        locations = loc_res.json()
        assert len(locations) > 0
        target_branch = locations[0]
        target_branch_id = target_branch["id"]
        target_branch_code = target_branch["code"]

        # 2. Prepare CSV for a unique product
        prod_name = "Catalog Reflected Test Beverage 500ml"
        prod_sku = "SKU-BEV-9999"
        prod_barcode = "8901122334455"
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow(["Product Name", "Category", "Barcode", "SKU", "Sale Price", "MRP", "Purchase Price", "Opening Stock", "Unit", "Branch Code / Name"])
        writer.writerow([prod_name, "Beverages", prod_barcode, prod_sku, "150.00", "175.00", "90.00", "42.0", "bottle", target_branch_code])
        csv_bytes = csv_buf.getvalue().encode("utf-8")

        files = {"file": ("beverages.csv", csv_bytes, "text/csv")}
        data = {
            "location_id": target_branch_id,
            "update_existing": "true",
            "create_categories": "true"
        }

        # 3. Import product
        import_res = await client.post("/api/v1/items/import", headers=headers, files=files, data=data)
        assert import_res.status_code == 200, import_res.text
        body = import_res.json()
        assert body["success"] is True
        assert body["createdCount"] + body["updatedCount"] >= 1

        # 4. Verify item is returned when querying with locationId (as done by InventoryView & PosBillingView)
        list_res = await client.get(f"/api/v1/items?locationId={target_branch_id}&search=Catalog+Reflected", headers=headers)
        assert list_res.status_code == 200
        catalog_items = list_res.json()["data"]
        assert len(catalog_items) >= 1
        matched_item = next((i for i in catalog_items if i["name"] == prod_name), None)
        assert matched_item is not None, f"Product '{prod_name}' was not returned in catalog for locationId '{target_branch_id}'"
        assert matched_item["sku"] == prod_sku

        # Verify location entry inside item
        item_loc_entry = next((l for l in matched_item["locations"] if l["locationId"] == target_branch_id), None)
        assert item_loc_entry is not None
        assert item_loc_entry["isListed"] is True
        assert float(item_loc_entry["currentStock"]) == 42.0

        # 5. Verify inventory movement is tracked with the real branch location
        move_res = await client.get(f"/api/v1/inventory/movements?locationId={target_branch_id}&search=Catalog+Reflected", headers=headers)
        assert move_res.status_code == 200
        movements = move_res.json()["data"]
        assert len(movements) >= 1
        matched_move = next((m for m in movements if m["sku"] == prod_sku or prod_name in (m.get("itemName") or "")), None)
        assert matched_move is not None
        assert matched_move["locationId"] == target_branch_id

